import os
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')

    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)


def setup_document():
    doc = Document()
    
    # Setup base font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Document styling for Headings
    for i in range(1, 4):
        h_style = doc.styles[f'Heading {i}']
        h_font = h_style.font
        h_font.name = 'Times New Roman'
        h_font.bold = True
        h_font.color.rgb = None # default black
        if i == 1:
            h_font.size = Pt(16)
        elif i == 2:
            h_font.size = Pt(14)
        else:
            h_font.size = Pt(12)

    return doc

def set_margins(section):
    section.left_margin = Cm(3.5)
    section.top_margin = Cm(2.5)
    section.right_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)

def format_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5

def add_placeholder_box(doc, caption_text, description_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("┌────────────────────────────────────────────┐\n│                                            │\n│       [INSERT IMAGE / DIAGRAM HERE]        │\n│                                            │\n└────────────────────────────────────────────┘")
    run.font.name = 'Courier New'
    run.font.size = Pt(10)
    
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(caption_text)
    cap_run.bold = True
    
    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc_run = desc.add_run(description_text)
    desc_run.italic = True
    format_paragraph(desc)

def create_report():
    doc = setup_document()
    
    # ------------------ COVER PAGE ------------------
    section = doc.sections[0]
    set_margins(section)
    
    doc.add_paragraph("\n\n\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("INDUSTRIAL TRAINING REPORT")
    run.bold = True
    run.font.size = Pt(18)
    
    doc.add_paragraph("\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ALFA: Automated Lead & Financial Analysis")
    run.bold = True
    run.font.size = Pt(16)
    
    doc.add_paragraph("\n\n")
    add_placeholder_box(doc, "", "[ABES LOGO / COLLEGE LOGO PLACEHOLDER]")
    doc.add_paragraph("\n\n")
    
    p = doc.add_paragraph("Submitted By")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    
    p = doc.add_paragraph("Name: Apurv Sharma\nUniversity Roll No.: 2300321540045")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("\n\n")
    p = doc.add_paragraph("Submitted To")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    
    p = doc.add_paragraph("Department of CSE-DS\nABES ENGINEERING COLLEGE\nGHAZIABAD")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    
    doc.add_page_break()

    # ------------------ CERTIFICATE ------------------
    p = doc.add_paragraph("CERTIFICATE BY COMPANY")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    
    doc.add_paragraph("\n\n")
    p = doc.add_paragraph("[INSERT ORIGINAL INTERNSHIP CERTIFICATE HERE]")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    
    add_placeholder_box(doc, "Figure: Internship Training Certificate", "Please replace this placeholder with the scanned copy or high-quality image of the official internship certificate provided by the company.")
    doc.add_page_break()

    # ------------------ DECLARATION ------------------
    p = doc.add_paragraph("DECLARATION")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    
    p = doc.add_paragraph("I hereby declare that the Industrial Training Report entitled \"ALFA: Automated Lead & Financial Analysis\" is an authentic record of my own work carried out during the Industrial Training period from 29/6/26 to 21/8/26 under the guidance of [ENTER FACULTY GUIDE NAME] (Faculty Guide) and Mr. Abhishek Singh (Company Mentor).")
    format_paragraph(p)
    p = doc.add_paragraph("The matter embodied in this report has not been submitted by me for the award of any other degree or diploma to any other University or Institution.")
    format_paragraph(p)
    
    doc.add_paragraph("\n\n\n\n\n")
    p = doc.add_paragraph("Apurv Sharma\n2300321540045")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_page_break()
    
    # ------------------ ACKNOWLEDGEMENT ------------------
    p = doc.add_paragraph("ACKNOWLEDGEMENT")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    
    p = doc.add_paragraph("I would like to express my deepest gratitude to all those who provided support, guidance, and encouragement throughout my industrial training. This project, \"ALFA: Automated Lead & Financial Analysis,\" would not have been possible without their invaluable contributions.")
    format_paragraph(p)
    
    p = doc.add_paragraph("First and foremost, I extend my sincere thanks to my company mentor, Mr. Abhishek Singh, for providing me with the opportunity to work on this challenging and highly impactful project at EY. His technical insights, strategic direction, and constant feedback were instrumental in the successful completion of the ALFA system.")
    format_paragraph(p)

    p = doc.add_paragraph("I am also profoundly grateful to my faculty guide, [ENTER FACULTY GUIDE NAME], for their continuous academic support, supervision, and motivation during the course of this training. Their guidance helped bridge the gap between theoretical concepts and practical industry implementation.")
    format_paragraph(p)

    p = doc.add_paragraph("I would like to express my sincere appreciation to the Department of CSE-DS and the management of ABES Engineering College, Ghaziabad, for providing a conducive environment and robust academic framework that prepared me for this industrial exposure.")
    format_paragraph(p)

    p = doc.add_paragraph("Lastly, I would like to thank my co-intern and project partner, Harshiv, for his collaborative spirit, dedication, and immense contribution to the development of the ALFA platform. Working together on complex Transfer Pricing rules and generative AI integrations was a tremendous learning experience.")
    format_paragraph(p)
    doc.add_page_break()
    
    # ------------------ ABOUT COMPANY ------------------
    p = doc.add_paragraph("ABOUT THE COMPANY")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    
    p = doc.add_paragraph("Ernst & Young (EY) is one of the largest multinational professional services networks in the world. It primarily provides assurance (including financial audit), tax, consulting, and advisory services to its clients. As one of the \"Big Four\" accounting firms, EY serves a massive portfolio of enterprise clients, facilitating compliance, risk management, and strategic business transformation.")
    format_paragraph(p)
    
    p = doc.add_paragraph("During this internship, the project was conducted within the Transfer Pricing (TP) domain. Transfer Pricing relates to the rules and methods for pricing transactions within and between enterprises under common ownership or control. Because of the complex regulatory requirements globally, TP teams must meticulously analyze Annual Reports (ARs) to prepare compliance documentation and Business Development (BD) Fact Sheets.")
    format_paragraph(p)

    p = doc.add_paragraph("The internship focused on applying modern software engineering paradigms—specifically Generative AI (Large Language Models), unstructured data extraction, and process automation—to solve core operational bottlenecks in the TP department. By automating the ingestion of massive financial documents, the firm can allocate human capital to higher-order analytical tasks rather than repetitive data entry.")
    format_paragraph(p)
    doc.add_page_break()

    # ------------------ TOC, TABLES, FIGURES (Placeholders for Word to generate) ------------------
    # We will just write placeholders because python-docx cannot easily generate real dynamic TOCs
    p = doc.add_paragraph("TABLE OF CONTENTS")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    p = doc.add_paragraph("[Please use Microsoft Word's 'References > Table of Contents' feature to generate the TOC here.]")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    
    p = doc.add_paragraph("LIST OF TABLES")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    p = doc.add_paragraph("[Please use Microsoft Word to generate the List of Tables here.]")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    p = doc.add_paragraph("LIST OF FIGURES")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    p = doc.add_paragraph("[Please use Microsoft Word to generate the List of Figures here.]")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    
    p = doc.add_paragraph("ABBREVIATIONS AND NOMENCLATURE")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    
    abbrevs = [
        ("AE", "Associated Enterprise"),
        ("AI", "Artificial Intelligence"),
        ("API", "Application Programming Interface"),
        ("AR", "Annual Report"),
        ("BD", "Business Development"),
        ("LLM", "Large Language Model"),
        ("OCR", "Optical Character Recognition"),
        ("PBT", "Profit Before Tax"),
        ("RPT", "Related Party Transactions"),
        ("TP", "Transfer Pricing")
    ]
    for abv, full in abbrevs:
        p = doc.add_paragraph(f"{abv} \t– {full}")
        format_paragraph(p)
    doc.add_page_break()

    # ------------------ CHAPTER 1 ------------------
    doc.add_heading('Chapter 1 – Introduction to Project', level=1)
    
    doc.add_heading('1.1 Introduction', level=2)
    p = doc.add_paragraph("The ALFA (Automated Lead & Financial Analysis) project is an intelligent document processing and financial analysis tool designed to streamline the operations of the Transfer Pricing (TP) department. Its primary function is to automate the extraction, analysis, and synthesis of complex corporate data from lengthy Annual Reports (ARs) and Financial Statement PDFs directly into standardized TP Business Development (BD) Fact Sheet Excel workbooks.")
    format_paragraph(p)
    p = doc.add_paragraph("Traditionally, financial analysts are required to manually read through annual reports—often exceeding 100 to 200 pages—to locate nuanced tax and financial data points. This manual process is tedious, time-consuming, and highly susceptible to human error. ALFA resolves this by employing state-of-the-art Generative AI techniques, dual-mode text and vision extraction, and strictly validated JSON schemas to accurately capture and format data without human intervention.")
    format_paragraph(p)

    doc.add_heading('1.2 Problem Statement', level=2)
    p = doc.add_paragraph("The preparation of TP Business Development Fact Sheets is a mandatory, repetitive task that consumes significant professional hours. The core challenges in the existing manual process are:")
    format_paragraph(p)
    
    p = doc.add_paragraph("1. Unstructured Data: Financial statements and Annual Reports lack a uniform structural layout across different companies, making standard rule-based scraping impossible.")
    p = doc.add_paragraph("2. Token Limits & Cost: Feeding entire 100+ page PDFs to Large Language Models (LLMs) exceeds context windows and incurs massive API costs.")
    p = doc.add_paragraph("3. Complex Extraction Rules: Extracting 'Turnover' requires separating 'Revenue from Operations' from 'Other Income'. Identifying 'Related Party Transactions (RPT)' requires filtering out natural persons (individuals/directors) and consolidating corporate entities. Traditional models fail to apply these domain-specific TP rules.")
    p = doc.add_paragraph("4. Non-Destructive Formatting: Transfer Pricing teams utilize highly specific, styled Excel templates. Automated tools often overwrite formatting, destroy formulas, and break merged cells when inserting new data.")
    
    doc.add_heading('1.3 Objectives', level=2)
    p = doc.add_paragraph("The primary objectives of the ALFA project were to:")
    format_paragraph(p)
    p = doc.add_paragraph("• Design a Dual-Mode Extraction Engine (Text + Vision) to dynamically parse financial PDFs and isolate only the relevant pages, drastically reducing token consumption.")
    p = doc.add_paragraph("• Implement a highly resilient AI pipeline using Google Gemini API and Pydantic v2 schemas to ensure the extracted data strictly conforms to Transfer Pricing logic.")
    p = doc.add_paragraph("• Develop a sophisticated, non-destructive Excel manipulation engine (Safe Sheet) that dynamically expands rows, maintains styles, and rewrites formulas without degrading the template's integrity.")
    p = doc.add_paragraph("• Build an enterprise-grade Streamlit web interface for human-in-the-loop review, enabling professionals to inspect, edit, and validate AI-generated insights before exporting final reports.")

    doc.add_heading('1.4 Scope of the Project', level=2)
    p = doc.add_paragraph("The scope of the ALFA system is restricted to the analysis of Annual Reports and Financial Statement PDFs for corporate entities. It aims to populate critical Fact Sheet fields including:")
    format_paragraph(p)
    p = doc.add_paragraph("• Group and entity headquarters and operational descriptions.")
    p = doc.add_paragraph("• Standalone and Consolidated financial summaries (Turnover, Total Costs).")
    p = doc.add_paragraph("• Associated Enterprise (AE) / Domestic / Export revenue splits based on RPTs.")
    p = doc.add_paragraph("• Director and Promoter Shareholding patterns.")
    p = doc.add_paragraph("• High-level Related Party Transactions (RPTs) filtered strictly for corporate entities.")
    p = doc.add_paragraph("• Litigation details extracted exclusively from Contingent Liabilities disclosures.")
    
    doc.add_heading('1.5 Project Workflow', level=2)
    add_placeholder_box(doc, "Figure 1.1 – ALFA System Workflow", "Insert a flowchart showing: PDF Upload -> Dual-Mode Extraction (pdfplumber/Gemini Vision) -> Pydantic Validation -> Streamlit UI Preview -> Excel Safe Sheet Engine -> Final ZIP Download.")
    
    p = doc.add_paragraph("The typical workflow starts with the user uploading the BD Fact Sheet template and multiple Annual Report PDFs via the Streamlit interface. The Orchestrator distributes the workload across a ThreadPoolExecutor. The PDF Processor filters the document using keyword scoring, routing either to native text extraction or Gemini Vision Fallback. The LLM extracts data into a strict JSON schema, which is parsed by Pydantic. The user reviews the data via interactive tables. Finally, the Excel Engine safely populates the template, injecting source-citing comments for auditability.")
    format_paragraph(p)

    doc.add_page_break()

    # ------------------ CHAPTER 2 ------------------
    doc.add_heading('Chapter 2 – Tools & Technology Used', level=1)
    
    p = doc.add_paragraph("ALFA was developed using a modern, scalable technology stack optimized for rapid prototyping, robust data validation, and seamless deployment. The primary technologies utilized in the project are outlined below.")
    format_paragraph(p)

    doc.add_heading('2.1 Programming Language: Python 3.10+', level=2)
    p = doc.add_paragraph("Python was chosen as the core language due to its unparalleled ecosystem for data manipulation, AI integration, and scripting. Features like type hinting and asynchronous support were heavily utilized to build resilient backend services.")
    format_paragraph(p)

    doc.add_heading('2.2 Artificial Intelligence & LLMs: Google Gemini', level=2)
    p = doc.add_paragraph("The core extraction intelligence is powered by Google's Gemini Multimodal models. Gemini was selected over competitors due to its massive context window and native Multimodal File API, which allows direct ingestion of images and scanned PDFs. The system utilizes a Dual-Mode engine:")
    format_paragraph(p)
    p = doc.add_paragraph("• Primary (Text) Mode: Uses pdfplumber to extract text from natively digital PDFs, isolating the most relevant ~40 pages based on financial keywords, reducing token usage by up to 80%.")
    p = doc.add_paragraph("• Fallback (Vision) Mode: Automatically triggered if heuristic detection finds fewer than 2,000 characters (indicating a scanned document), routing the entire file to the Gemini Vision API.")

    doc.add_heading('2.3 Frontend & UI Framework: Streamlit', level=2)
    p = doc.add_paragraph("Streamlit was utilized to rapidly develop the enterprise web interface. It provided out-of-the-box components for file uploads, metric cards, and session state management. The st.data_editor component was critically important, allowing the implementation of a 'Human-in-the-loop' workflow where analysts can review and correct the AI's extracted financial tables before triggering the Excel export.")
    format_paragraph(p)

    doc.add_heading('2.4 Data Validation: Pydantic v2', level=2)
    p = doc.add_paragraph("Pydantic is a data validation library for Python. In ALFA, Pydantic v2 was used to define strict JSON schemas for the LLM output. Field validators were implemented to enforce numeric coercion, reject empty names, and gracefully handle null values. This ensures the application never crashes due to malformed AI hallucinations.")
    format_paragraph(p)

    doc.add_heading('2.5 Excel Manipulation: Openpyxl', level=2)
    p = doc.add_paragraph("Openpyxl was the engine behind the non-destructive Excel modification. Traditional tools like Pandas overwrite complex Excel styles. Openpyxl allowed the creation of a 'Safe Sheet' module which preserves native fonts, cell colors, borders, and number formats. It enabled dynamic row expansion—shifting formulas safely via regex, rewriting merged cells, and cloning styles dynamically.")
    format_paragraph(p)

    doc.add_heading('2.6 Resilience & Concurrency: Tenacity & ThreadPoolExecutor', level=2)
    p = doc.add_paragraph("To handle API rate limits (e.g., HTTP 429 or 503 errors) from the LLM provider, the 'tenacity' library was employed to implement exponential-backoff retries (up to 5 attempts, with 2-30s wait times). For performance optimization, Python's ThreadPoolExecutor allowed concurrent batch processing of multiple PDFs (1-5 workers).")
    format_paragraph(p)

    doc.add_heading('2.7 Containerization: Docker', level=2)
    p = doc.add_paragraph("To ensure cross-platform compatibility and ease of deployment, the application was containerized using Docker. A multi-stage Dockerfile and a docker-compose.yml configuration were created, encapsulating all dependencies, environment variables, and execution environments.")
    format_paragraph(p)

    doc.add_page_break()

    # ------------------ CHAPTER 3 ------------------
    doc.add_heading('Chapter 3 – Snapshots', level=1)
    
    p = doc.add_paragraph("This chapter presents the visual interface and execution flow of the ALFA platform through various screenshots captured during the operational testing of the application.")
    format_paragraph(p)

    add_placeholder_box(doc, "Figure 3.1 – Main Application Dashboard", "Insert a screenshot showing the main Streamlit UI, including the sidebar for API key input, the file upload zones for the Template and Annual Reports, and the metric cards showing processing status.")
    
    p = doc.add_paragraph("Figure 3.1 illustrates the primary user interface. The sidebar provides configuration inputs, while the main area is dedicated to file uploads. The system utilizes session state to retain the uploaded documents across reruns, ensuring a seamless user experience.")
    format_paragraph(p)

    add_placeholder_box(doc, "Figure 3.2 – Human-in-the-Loop Data Editor", "Insert a screenshot showing the st.data_editor tables where the extracted Shareholding, RPT, and Litigation data is presented in an editable grid format.")
    
    p = doc.add_paragraph("Figure 3.2 displays the human-in-the-loop preview mechanism. Before the data is committed to the Excel template, the user is presented with interactive data grids. They can modify extracted values, correct hallucinations, or adjust TP classifications directly in the UI, adhering to the principle that AI outputs are drafts requiring human validation.")
    format_paragraph(p)

    add_placeholder_box(doc, "Figure 3.3 – Populated Excel Output with Source Comments", "Insert a screenshot of the final Excel file, highlighting a cell that has an Excel Comment (red triangle) showing the source citation extracted by the LLM.")
    
    p = doc.add_paragraph("Figure 3.3 demonstrates the final output generated by the openpyxl engine. The formatting of the BD Fact Sheet template is completely preserved. Notably, every populated cell includes an Excel comment citing the exact page or section from the Annual Report where the data was sourced, enabling rapid auditability by reviewers.")
    format_paragraph(p)

    doc.add_page_break()

    # ------------------ CHAPTER 4 ------------------
    doc.add_heading('Chapter 4 – Results and Discussions', level=1)
    
    doc.add_heading('4.1 System Testing Methodology', level=2)
    p = doc.add_paragraph("The ALFA system underwent rigorous module-wise testing to validate its accuracy against the strict domain rules defined in the SKILL.md specification. Since standard programmatic unit tests cannot entirely evaluate generative AI extraction accuracy, a hybrid methodology was adopted:")
    format_paragraph(p)
    p = doc.add_paragraph("1. Schema Validation Tests: Automated Pytest scripts verified that the Pydantic models correctly rejected malformed data and coerced data types appropriately.")
    p = doc.add_paragraph("2. Formula Shift Testing: Isolated testing of the 'safe_sheet.py' engine to ensure that inserting rows dynamically updated regex-matched Excel formulas without corrupting the workbook.")
    p = doc.add_paragraph("3. LLM Output Validation: Manual benchmarking of the AI's output against a set of historical Annual Reports to verify adherence to Transfer Pricing rules (e.g., excluding non-corporate entities from RPTs, extracting Turnover strictly as Revenue from Operations).")

    doc.add_heading('4.2 Functional Results', level=2)
    p = doc.add_paragraph("The implementation successfully automated the Fact Sheet population process. The Dual-Mode extraction engine proved highly effective. By utilizing pdfplumber for keyword-scored page filtering, the system consistently reduced the payload sent to the LLM by 70-80%, mitigating token limit errors and significantly reducing API costs.")
    format_paragraph(p)
    
    p = doc.add_paragraph("The non-destructive Excel engine performed flawlessly. It successfully executed dynamic row expansion, preserving merged cells and cloning styles. The automated insertion of source-citing cell comments drastically reduced the time required for human review.")
    format_paragraph(p)

    doc.add_heading('4.3 Performance Observations', level=2)
    p = doc.add_paragraph("Performance metrics were not formally recorded during the training period as a production SLA benchmark; however, qualitative observations indicate significant time savings. Processing a 150-page Annual Report manually typically takes a financial analyst several hours. The ALFA system processes and synthesizes the same document into the Excel template in under 3 minutes, representing an enormous increase in operational efficiency.")
    format_paragraph(p)
    
    doc.add_heading('4.4 Challenges Encountered & Resolutions', level=2)
    p = doc.add_paragraph("Challenge: LLMs struggling to differentiate between natural persons (Directors) and corporate entities in Related Party Transaction tables.")
    p = doc.add_paragraph("Resolution: Highly explicit prompt engineering instructions were added to the SKILL.md rules, strictly commanding the model to ignore individuals. Additionally, human-in-the-loop editing was introduced as a fail-safe.")
    
    p = doc.add_paragraph("Challenge: API Rate Limiting (429 Too Many Requests).")
    p = doc.add_paragraph("Resolution: The 'tenacity' library was integrated to handle exponential backoff, ensuring the concurrent ThreadPoolExecutor did not overwhelm the API endpoint.")

    doc.add_page_break()

    # ------------------ CHAPTER 5 ------------------
    doc.add_heading('Chapter 5 – Conclusions and Future Scope', level=1)
    
    doc.add_heading('5.1 Conclusion', level=2)
    p = doc.add_paragraph("The ALFA project successfully demonstrated the viability of applying advanced Generative AI and unstructured data processing techniques to highly regulated financial tasks. During this industrial training at EY, I successfully contributed to a robust automation pipeline that mitigates the manual burden of Transfer Pricing BD Fact Sheet generation. The architecture is resilient, cost-effective, and maintains the strict formatting requirements of enterprise Excel templates. The system not only automates data entry but fundamentally transforms the analyst workflow from manual extraction to strategic review and validation.")
    format_paragraph(p)

    doc.add_heading('5.2 Learning Outcomes', level=2)
    p = doc.add_paragraph("This internship provided profound practical experience in modern software engineering and enterprise AI deployment. Key learnings included:")
    format_paragraph(p)
    p = doc.add_paragraph("• Designing fault-tolerant AI applications using structured schemas (Pydantic).")
    p = doc.add_paragraph("• Manipulating complex Excel DOMs programmatically without destructive consequences.")
    p = doc.add_paragraph("• Understanding the nuances of Prompt Engineering for highly specific financial compliance rules.")
    p = doc.add_paragraph("• Working with asynchronous execution, ThreadPool processing, and containerization (Docker).")

    doc.add_heading('5.3 Future Scope', level=2)
    p = doc.add_paragraph("While the current version of ALFA is highly functional, several enhancements can be developed in the future:")
    format_paragraph(p)
    p = doc.add_paragraph("• Direct ERP Integration: Bypassing the PDF analysis entirely for internal clients by integrating directly with ERP databases via secured APIs.")
    p = doc.add_paragraph("• Advanced RAG (Retrieval-Augmented Generation): Implementing vector databases to allow analysts to 'chat' with the financial document natively within the Streamlit app for custom queries.")
    p = doc.add_paragraph("• Multi-Language Support: Adapting the extraction models to process Annual Reports from non-English jurisdictions, expanding the tool's global applicability.")

    doc.add_page_break()

    # ------------------ REFERENCES ------------------
    doc.add_heading('References', level=1)
    
    refs = [
        "Python Software Foundation. (2024). Python 3.10 Documentation. Retrieved from https://docs.python.org/3.10/",
        "Google Cloud. (2024). Google Gemini API Documentation. Retrieved from https://ai.google.dev/",
        "Streamlit Inc. (2024). Streamlit Documentation. Retrieved from https://docs.streamlit.io/",
        "Pydantic. (2024). Pydantic V2 Documentation. Retrieved from https://docs.pydantic.dev/",
        "Openpyxl Project. (2024). openpyxl - A Python library to read/write Excel 2010 xlsx/xlsm files. Retrieved from https://openpyxl.readthedocs.io/",
        "pdfplumber Project. (2024). pdfplumber Documentation. Retrieved from https://github.com/jsvine/pdfplumber"
    ]
    
    for r in refs:
        p = doc.add_paragraph(r)
        format_paragraph(p)
        
    doc.add_page_break()
    
    # ------------------ PLACEHOLDER CHECKLIST ------------------
    doc.add_heading('PLACEHOLDER CHECKLIST', level=1)
    p = doc.add_paragraph("The following items require manual review and replacement before the final submission:")
    format_paragraph(p)
    
    p = doc.add_paragraph("1. [ABES LOGO / COLLEGE LOGO PLACEHOLDER] - Cover Page")
    p = doc.add_paragraph("2. [INSERT ORIGINAL INTERNSHIP CERTIFICATE HERE] - Certificate Page")
    p = doc.add_paragraph("3. [ENTER FACULTY GUIDE NAME] - Declaration and Acknowledgement Pages")
    p = doc.add_paragraph("4. Table of Contents - Needs to be generated in Word via References tab")
    p = doc.add_paragraph("5. List of Tables - Needs to be generated in Word")
    p = doc.add_paragraph("6. List of Figures - Needs to be generated in Word")
    p = doc.add_paragraph("7. Figure 1.1 - Project Workflow Diagram Image")
    p = doc.add_paragraph("8. Figure 3.1 - Main Application Dashboard Image")
    p = doc.add_paragraph("9. Figure 3.2 - Human-in-the-Loop Data Editor Image")
    p = doc.add_paragraph("10. Figure 3.3 - Populated Excel Output Image")
    
    doc.save('Industrial_Training_Report_Apurv_Sharma_ALFA.docx')

if __name__ == "__main__":
    create_report()
