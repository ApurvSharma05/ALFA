"""
tests/test_formula_shift.py

Tests for the safe_sheet module: formula rewriting and end-to-end row insertion.
"""

import os
import tempfile
import pytest
from openpyxl import Workbook, load_workbook

from core.safe_sheet import shift_formula, insert_rows_safe


# ---------------------------------------------------------------------------
# shift_formula unit tests
# ---------------------------------------------------------------------------

class TestShiftFormula:
    def test_simple_reference(self):
        assert shift_formula("=D12-D13", threshold=13, delta=2) == "=D12-D15"

    def test_both_refs_shifted(self):
        assert shift_formula("=D12-D13", threshold=12, delta=3) == "=D15-D16"

    def test_sum_range(self):
        assert shift_formula("=SUM(D42:D43)", threshold=43, delta=2) == "=SUM(D42:D45)"

    def test_absolute_row(self):
        # $44 should still shift because the row physically moved
        assert shift_formula("=D42/$D$44", threshold=43, delta=2) == "=D42/$D$46"

    def test_absolute_col(self):
        assert shift_formula("=$A10", threshold=10, delta=1) == "=$A11"

    def test_no_shift_below_threshold(self):
        assert shift_formula("=A1+B2", threshold=10, delta=5) == "=A1+B2"

    def test_mixed_shift(self):
        result = shift_formula("=A5+A15", threshold=10, delta=3)
        assert result == "=A5+A18"

    def test_multi_column_refs(self):
        result = shift_formula("=SUM(AA100:AB200)", threshold=150, delta=10)
        assert result == "=SUM(AA100:AB210)"

    def test_complex_formula(self):
        result = shift_formula("=IF(D43>0,D43/D44,0)", threshold=43, delta=2)
        assert result == "=IF(D45>0,D45/D46,0)"


# ---------------------------------------------------------------------------
# insert_rows_safe integration tests
# ---------------------------------------------------------------------------

class TestInsertRowsSafe:
    def _create_test_workbook(self, path):
        """Create a minimal workbook that mimics the shareholding table."""
        wb = Workbook()
        ws = wb.active
        ws.title = "TestSheet"

        # Some fixed fields above
        ws["A1"] = "Header"
        ws["C3"] = "HQ Address"

        # Shareholding table
        ws["C40"] = "Shareholding Header"
        ws["C42"] = "Director A"
        ws["D42"] = 5000
        ws["C43"] = "Total"
        ws["D43"] = "=SUM(D42:D42)"

        # Row below that references the table
        ws["C45"] = "RPT Section"
        ws["D45"] = "=D43*2"

        # Merge a range below
        ws.merge_cells("B47:B50")

        wb.save(path)
        return path

    def test_delta_zero_is_noop(self):
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "src.xlsx")
            out = os.path.join(td, "out.xlsx")
            self._create_test_workbook(src)

            insert_rows_safe(src, out, "TestSheet", threshold=43, delta=0)

            wb = load_workbook(out)
            ws = wb["TestSheet"]
            assert ws["C43"].value == "Total"
            assert ws["D43"].value == "=SUM(D42:D42)"

    def test_insert_2_rows(self):
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "src.xlsx")
            out = os.path.join(td, "out.xlsx")
            self._create_test_workbook(src)

            insert_rows_safe(src, out, "TestSheet", threshold=43, delta=2)

            wb = load_workbook(out)
            ws = wb["TestSheet"]

            # Fixed fields should not move
            assert ws["A1"].value == "Header"
            assert ws["C3"].value == "HQ Address"
            assert ws["C42"].value == "Director A"

            # Total row should have shifted from 43 → 45
            assert ws["C45"].value == "Total"
            # The formula =SUM(D42:D42) — both refs are row 42 (< threshold 43),
            # so they don't shift. The formula text stays unchanged.
            assert ws["D45"].value == "=SUM(D42:D42)"

            # RPT section shifted from 45 → 47
            assert ws["C47"].value == "RPT Section"
            # =D43*2 had ref at row 43 (>= threshold), shifts to D45
            assert ws["D47"].value == "=D45*2"

    def test_insert_5_rows(self):
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "src.xlsx")
            out = os.path.join(td, "out.xlsx")
            self._create_test_workbook(src)

            insert_rows_safe(src, out, "TestSheet", threshold=43, delta=5)

            wb = load_workbook(out)
            ws = wb["TestSheet"]

            # Total row: 43 + 5 = 48
            assert ws["C48"].value == "Total"
            # Same as above: both D42 refs are below threshold 43
            assert ws["D48"].value == "=SUM(D42:D42)"

    def test_invalid_threshold_raises(self):
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "src.xlsx")
            out = os.path.join(td, "out.xlsx")
            self._create_test_workbook(src)

            with pytest.raises(ValueError, match="threshold must be >= 1"):
                insert_rows_safe(src, out, "TestSheet", threshold=0, delta=2)

    def test_invalid_delta_raises(self):
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "src.xlsx")
            out = os.path.join(td, "out.xlsx")
            self._create_test_workbook(src)

            with pytest.raises(ValueError, match="delta must be >= 0"):
                insert_rows_safe(src, out, "TestSheet", threshold=43, delta=-1)

    def test_missing_sheet_raises(self):
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "src.xlsx")
            out = os.path.join(td, "out.xlsx")
            self._create_test_workbook(src)

            with pytest.raises(ValueError, match="Sheet 'NoSuchSheet' not found"):
                insert_rows_safe(src, out, "NoSuchSheet", threshold=43, delta=2)

    def test_styles_preserved_on_new_rows(self):
        """New rows should inherit formatting from the row above the gap."""
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "src.xlsx")
            out = os.path.join(td, "out.xlsx")

            # Create workbook with styled row 42
            wb = Workbook()
            ws = wb.active
            ws.title = "TestSheet"
            from openpyxl.styles import Font
            ws["C42"] = "Data Row"
            ws["C42"].font = Font(bold=True, size=14)
            ws["C43"] = "Total"
            wb.save(src)

            insert_rows_safe(src, out, "TestSheet", threshold=43, delta=1)

            wb2 = load_workbook(out)
            ws2 = wb2["TestSheet"]
            # New row 43 should have the bold font from row 42
            assert ws2["C43"].font.bold is True
            assert ws2["C43"].font.size == 14
