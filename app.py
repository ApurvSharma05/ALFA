"""
ALFA (Automated Lead & Financial Analysis) — Streamlit Application

Thin orchestration layer that coordinates:
  - config/  → settings and template mappings
  - core/    → PDF processing, Gemini extraction, Excel population
  - ui/      → styled components, previews, downloads

Three-phase workflow:
  1. EXTRACT  — process PDFs via Gemini (concurrent)
  2. REVIEW   — editable preview tables for human-in-the-loop overrides
  3. GENERATE — write to Excel and offer downloads (individual + ZIP)
"""

import os
import json
import tempfile
import logging
from concurrent.futures import ThreadPoolExecutor, as_completed

import streamlit as st
from openpyxl import load_workbook

from config.settings import load_settings, get_system_prompt, logger
from config.template_mapping import FactSheetMapping
from core.pdf_processor import validate_pdf, extract_relevant_pages, is_scanned_pdf
from core.extractor import GeminiExtractor
from core.excel_engine import fill_factsheet
from core.safe_sheet import insert_rows_safe
from ui.styles import inject_styles
from ui.components import (
    render_metrics_bar,
    render_company_card,
    render_preview_editor,
    render_download_section,
)

# ---------------------------------------------------------------------------
# Page config (must be first Streamlit call)
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title="ALFA — Automated Lead & Financial Analysis",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

inject_styles()

# ---------------------------------------------------------------------------
# Session state initialization
# ---------------------------------------------------------------------------
_DEFAULTS = {
    "extraction_results": {},   # company_name → dict (extracted JSON)
    "extraction_errors": {},    # company_name → str (error message)
    "extraction_modes": {},     # company_name → "text" | "vision"
    "excel_bytes": {},          # company_name → bytes (populated xlsx)
    "fill_warnings": {},        # company_name → list[str]
    "phase": "upload",          # "upload" | "extract" | "review" | "generate"
    "processing_complete": False,
}
for key, default in _DEFAULTS.items():
    if key not in st.session_state:
        st.session_state[key] = default


# ---------------------------------------------------------------------------
# Sidebar
# ---------------------------------------------------------------------------
settings = load_settings()

with st.sidebar:
    st.header("⚙️ Settings")
    api_key = st.text_input(
        "Gemini API Key",
        value=settings.gemini_api_key,
        type="password",
        help="Get a key at https://aistudio.google.com/",
    )
    model_choice = st.text_input("Model", value=settings.gemini_model)
    max_workers = st.slider("Concurrency", min_value=1, max_value=5, value=settings.max_workers,
                            help="Number of PDFs to process simultaneously")

    with st.expander("🔧 Advanced"):
        vision_threshold = st.number_input(
            "Vision fallback threshold (chars)",
            value=settings.vision_fallback_threshold,
            min_value=500,
            max_value=10000,
            step=500,
        )
        max_pages = st.number_input(
            "Max relevant pages per PDF",
            value=settings.max_relevant_pages,
            min_value=10,
            max_value=200,
            step=5,
        )

    st.markdown("---")
    st.caption("ALFA v2.0 — Refactored Architecture")


# ---------------------------------------------------------------------------
# Header
# ---------------------------------------------------------------------------
st.title("📊 ALFA")
st.caption("Automated Lead & Financial Analysis — TP BD Fact Sheet Automation")

with st.expander("ℹ️ How to use ALFA"):
    st.markdown("""
**ALFA** automates the extraction of financial data from Annual Reports (ARs)
into standardized BD Fact Sheet Excel templates.

**Workflow:**
1. Upload your BD template (`.xlsx`) and one or more Annual Report PDFs
2. Click **Start Extraction** — ALFA processes PDFs concurrently via Gemini
3. **Review** extracted data in interactive tables — edit any values
4. Click **Generate Excel** — download individual files or a bulk ZIP
    """)


# ---------------------------------------------------------------------------
# File uploads
# ---------------------------------------------------------------------------
col1, col2 = st.columns(2)
with col1:
    excel_template = st.file_uploader("📁 Upload BD Template", type=["xlsx"])
with col2:
    pdf_files = st.file_uploader("📄 Upload Annual Reports", type=["pdf"], accept_multiple_files=True)


# ---------------------------------------------------------------------------
# Phase 1: EXTRACTION
# ---------------------------------------------------------------------------

def _process_single_pdf(pdf_path: str, company_name: str, extractor: GeminiExtractor,
                        vision_threshold: int, max_pages: int) -> dict:
    """Process one PDF: validate → extract text → call Gemini. Returns result dict."""
    info = validate_pdf(pdf_path)
    if not info.is_valid:
        return {"company": company_name, "error": f"Invalid PDF: {info.error}", "mode": "none"}

    text = extract_relevant_pages(pdf_path, max_pages=max_pages)
    if is_scanned_pdf(text, threshold=vision_threshold):
        result = extractor.extract_from_pdf(pdf_path, company_name)
    else:
        result = extractor.extract_from_text(text, company_name)

    if result.success:
        return {"company": company_name, "data": result.data, "mode": result.mode}
    else:
        return {"company": company_name, "error": result.error, "mode": result.mode}


if st.button("🚀 Start Extraction", type="primary", disabled=not (api_key and excel_template and pdf_files)):
    # Reset state
    for key in ["extraction_results", "extraction_errors", "extraction_modes", "excel_bytes", "fill_warnings"]:
        st.session_state[key] = {}
    st.session_state["phase"] = "extract"
    st.session_state["processing_complete"] = False

    extractor = GeminiExtractor(api_key=api_key, model=model_choice)
    total = len(pdf_files)
    progress_bar = st.progress(0, text="Initializing...")

    with tempfile.TemporaryDirectory() as temp_dir:
        # Save PDFs to temp
        pdf_paths = {}
        for pdf_file in pdf_files:
            name = pdf_file.name.replace(".pdf", "")
            path = os.path.join(temp_dir, f"{name}.pdf")
            with open(path, "wb") as f:
                f.write(pdf_file.getvalue())
            pdf_paths[name] = path

        # Concurrent extraction
        completed = 0
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = {
                executor.submit(
                    _process_single_pdf, path, name, extractor, vision_threshold, max_pages
                ): name
                for name, path in pdf_paths.items()
            }

            for future in as_completed(futures):
                company = futures[future]
                completed += 1
                progress_bar.progress(
                    completed / total,
                    text=f"Processed {completed}/{total}: {company}",
                )

                try:
                    result = future.result()
                    if "error" in result:
                        st.session_state["extraction_errors"][company] = result["error"]
                    else:
                        st.session_state["extraction_results"][company] = result["data"]
                    st.session_state["extraction_modes"][company] = result.get("mode", "unknown")
                except Exception as exc:
                    logger.error("[%s] Unexpected error: %s", company, exc, exc_info=True)
                    st.session_state["extraction_errors"][company] = str(exc)

    progress_bar.progress(1.0, text="Extraction complete!")
    st.session_state["phase"] = "review"
    st.session_state["processing_complete"] = True
    st.rerun()


# ---------------------------------------------------------------------------
# Metrics bar (shown after extraction)
# ---------------------------------------------------------------------------
if st.session_state["processing_complete"]:
    results = st.session_state["extraction_results"]
    errors = st.session_state["extraction_errors"]
    modes = st.session_state["extraction_modes"]
    total = len(results) + len(errors)

    render_metrics_bar(
        total=total,
        success=len(results),
        failed=len(errors),
        pending=0,
    )

    st.markdown("---")

    # Status cards
    for name in sorted(results.keys()):
        render_company_card(name, "success", mode=modes.get(name, ""))
    for name, err in sorted(errors.items()):
        render_company_card(name, "failed", error=err)


# ---------------------------------------------------------------------------
# Phase 2: REVIEW (editable previews)
# ---------------------------------------------------------------------------
if st.session_state["phase"] == "review" and st.session_state["extraction_results"]:
    st.markdown("---")
    st.header("📝 Review & Edit Extracted Data")
    st.caption("Modify any values below before generating the final Excel files.")

    for company_name, data in st.session_state["extraction_results"].items():
        with st.container():
            updated = render_preview_editor(data, company_name)
            st.session_state["extraction_results"][company_name] = updated

    # ---------------------------------------------------------------------------
    # Phase 3: GENERATE EXCEL
    # ---------------------------------------------------------------------------
    st.markdown("---")
    if st.button("📊 Generate Excel Files", type="primary"):
        if not excel_template:
            st.error("Template file is required. Please re-upload.")
        else:
            mapping = FactSheetMapping()
            gen_progress = st.progress(0, text="Generating Excel files...")

            with tempfile.TemporaryDirectory() as temp_dir:
                template_path = os.path.join(temp_dir, "template.xlsx")
                with open(template_path, "wb") as f:
                    f.write(excel_template.getvalue())

                wb_check = load_workbook(template_path)
                sheet_name = wb_check.sheetnames[0]
                wb_check.close()

                results = st.session_state["extraction_results"]
                total_gen = len(results)

                for idx, (company_name, data) in enumerate(results.items()):
                    try:
                        # Save JSON
                        json_path = os.path.join(temp_dir, f"{company_name}.json")
                        with open(json_path, "w") as f:
                            json.dump(data, f, indent=2)

                        # Expand shareholding if needed
                        template_to_use = template_path
                        required_sh_rows = len(data["fields"]["shareholding"]["rows"])
                        if required_sh_rows > mapping.shareholding_default_capacity:
                            expanded_path = os.path.join(temp_dir, f"expanded_{company_name}.xlsx")
                            delta = required_sh_rows - mapping.shareholding_default_capacity
                            insert_rows_safe(
                                src_path=template_path,
                                out_path=expanded_path,
                                sheet_name=sheet_name,
                                threshold=mapping.shareholding_insert_threshold,
                                delta=delta,
                            )
                            template_to_use = expanded_path

                        # Fill
                        out_path = os.path.join(temp_dir, f"{company_name}_Filled.xlsx")
                        fill_result = fill_factsheet(
                            data=data,
                            template_path=template_to_use,
                            out_path=out_path,
                            sheet_name=sheet_name,
                            mapping=mapping,
                        )

                        # Store warnings
                        if fill_result.warnings:
                            st.session_state["fill_warnings"][company_name] = fill_result.warnings
                            for w in fill_result.warnings:
                                st.warning(f"[{company_name}] {w}")

                        # Read bytes
                        with open(out_path, "rb") as f:
                            st.session_state["excel_bytes"][company_name] = f.read()

                    except Exception as exc:
                        logger.error("[%s] Excel generation failed: %s", company_name, exc, exc_info=True)
                        st.error(f"[{company_name}] Excel generation failed: {exc}")

                    gen_progress.progress((idx + 1) / total_gen, text=f"Generated {idx + 1}/{total_gen}")

            gen_progress.progress(1.0, text="All files generated!")
            st.session_state["phase"] = "generate"
            st.rerun()


# ---------------------------------------------------------------------------
# Phase 3: DOWNLOAD
# ---------------------------------------------------------------------------
if st.session_state["phase"] == "generate" and st.session_state["excel_bytes"]:
    render_download_section(
        results=st.session_state["excel_bytes"],
        errors=st.session_state["extraction_errors"],
    )


# ---------------------------------------------------------------------------
# Footer
# ---------------------------------------------------------------------------
st.markdown("---")
st.caption("⚠️ All outputs are drafts for human review. Verify extracted figures against the source Annual Report.")