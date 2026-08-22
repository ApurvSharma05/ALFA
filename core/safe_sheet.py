"""
core/safe_sheet.py

Non-destructive row insertion for openpyxl worksheets.

Refactored from safe_insert_rows.py with:
  - Input validation and boundary checks
  - Structured logging
  - Comprehensive docstrings
  - Same proven algorithm: snapshot → clear → shift → rewrite → re-merge

The core regex-based formula shifting is unchanged — it passed all tests in
the original implementation.
"""

import re
import logging
from copy import copy
from openpyxl import load_workbook

logger = logging.getLogger("alfa.safe_sheet")

# Matches A1-style cell references, respecting $ anchors.
# Groups: (col_abs, col_letters, row_abs, row_digits)
CELL_REF_RE = re.compile(r'(\$?)([A-Z]{1,3})(\$?)(\d+)')


def shift_formula(formula: str, threshold: int, delta: int) -> str:
    """
    Rewrite every A1-style cell reference in *formula* whose row is
    >= *threshold* by adding *delta* to the row number.

    Respects absolute row anchors ($): ``$A$10`` is shifted only if ``10 >= threshold``
    (the ``$`` simply means the user pinned it, but the row number still moves
    when we physically insert rows).

    Args:
        formula:   an Excel formula string starting with '='
        threshold: rows >= this value get shifted
        delta:     number of rows to shift down

    Returns:
        The formula with row numbers adjusted.
    """
    def _repl(m):
        col_abs, col, row_abs, row = m.groups()
        row_n = int(row)
        if row_n >= threshold:
            row_n += delta
        return f"{col_abs}{col}{row_abs}{row_n}"

    return CELL_REF_RE.sub(_repl, formula)


def _shift_ref(ref: str, threshold: int, delta: int) -> str:
    """Shift a merged-cell range string like 'B40:B45'."""
    return shift_formula(ref, threshold, delta)


def insert_rows_safe(
    src_path: str,
    out_path: str,
    sheet_name: str,
    threshold: int,
    delta: int,
) -> str:
    """
    Insert *delta* blank rows at *threshold* in *sheet_name*, preserving
    all cell values, formulas, styles, merged ranges, and row heights.

    **Algorithm:**

    1. Snapshot every cell (value, formula, font, fill, border, alignment,
       number format, protection) — including blank cells with formatting.
    2. Snapshot merged-cell ranges and row heights.
    3. Clear the sheet.
    4. Re-write each cell at its new position (old_row + delta if >= threshold).
    5. Shift formula references inside each formula string.
    6. Copy formatting from the row above the gap into the new blank rows so
       they inherit the data-row style rather than the total-row style.
    7. Re-create merges and row heights at their shifted positions.

    Args:
        src_path:   input Excel file
        out_path:   output Excel file (can be same as src_path)
        sheet_name: target worksheet name
        threshold:  row number at which to insert (1-based)
        delta:      number of rows to insert

    Returns:
        out_path

    Raises:
        ValueError: if threshold < 1 or delta < 0 or sheet not found
    """
    if threshold < 1:
        raise ValueError(f"threshold must be >= 1, got {threshold}")
    if delta < 0:
        raise ValueError(f"delta must be >= 0, got {delta}")

    if delta == 0:
        logger.debug("delta=0, nothing to insert — copying file as-is")
        wb = load_workbook(src_path)
        wb.save(out_path)
        return out_path

    wb = load_workbook(src_path, data_only=False)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")

    ws = wb[sheet_name]
    max_row = ws.max_row
    max_col = ws.max_column

    logger.info(
        "Inserting %d rows at threshold=%d in '%s' (max_row=%d, max_col=%d)",
        delta, threshold, sheet_name, max_row, max_col,
    )

    # 1. Snapshot merges
    old_merges = [str(r) for r in ws.merged_cells.ranges]
    for r in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(r))

    # 2. Snapshot row heights
    old_row_heights = {
        r: dim.height
        for r, dim in ws.row_dimensions.items()
        if dim.height is not None
    }

    # 3. Snapshot every cell
    snapshot = {}
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            snapshot[(cell.row, cell.column)] = {
                "value": cell.value,
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "alignment": copy(cell.alignment),
                "number_format": cell.number_format,
                "protection": copy(cell.protection),
            }

    # 4. Clear all cells (including the expanded area)
    for row in ws.iter_rows(min_row=1, max_row=max_row + delta + 5, max_col=max_col):
        for cell in row:
            cell.value = None

    # 5. Re-write each cell at its new position
    for (r, c), info in snapshot.items():
        new_r = r + delta if r >= threshold else r
        cell = ws.cell(row=new_r, column=c)
        val = info["value"]
        if isinstance(val, str) and val.startswith("="):
            val = shift_formula(val, threshold, delta)
        cell.value = val
        cell.font = info["font"]
        cell.fill = info["fill"]
        cell.border = info["border"]
        cell.alignment = info["alignment"]
        cell.number_format = info["number_format"]
        cell.protection = info["protection"]

    # 6. Copy styles from the row above the gap into the new blank rows
    style_source_row = threshold - 1
    for new_r in range(threshold, threshold + delta):
        for c in range(1, max_col + 1):
            src_info = snapshot.get((style_source_row, c))
            if src_info:
                tgt_cell = ws.cell(row=new_r, column=c)
                tgt_cell.font = copy(src_info["font"])
                tgt_cell.fill = copy(src_info["fill"])
                tgt_cell.border = copy(src_info["border"])
                tgt_cell.alignment = copy(src_info["alignment"])
                tgt_cell.number_format = src_info["number_format"]

    # 7. Re-create row heights and merges at their shifted positions
    for r, height in old_row_heights.items():
        new_r = r + delta if r >= threshold else r
        ws.row_dimensions[new_r].height = height

    for m in old_merges:
        new_m = _shift_ref(m, threshold, delta)
        ws.merge_cells(new_m)

    wb.save(out_path)
    logger.info("Saved expanded workbook to %s", out_path)
    return out_path
